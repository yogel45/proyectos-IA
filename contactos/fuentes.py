"""De donde sale el directorio.

Nadie teclea fichas: el directorio se construye con lo que la oficina ya produce.

  * **Nomina** (Excel del biometrico): el personal del despacho y su area.
  * **Central telefonica** (Excel de RingCentral): con quien se habla de verdad,
    cuanto y quien atendio. Un numero que solo llamo una vez y colgo no merece
    entrar al directorio; uno que llama seis veces, si.
  * **Documentos clasificados** (base de datos de DocuFlow AI): las partes, el
    despacho que firma, la aseguradora, el medico... y el papel de cada uno,
    deducido de la categoria del documento en el que aparece.
"""
from __future__ import annotations

import json
import logging
import re
import sqlite3
from collections import defaultdict
from datetime import date, datetime, time as dtime
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from . import db, modelo
from .config import CONFIG, DOCS_DB, SAMPLE_DIR

log = logging.getLogger("contactos.fuentes")

# Papel de cada persona segun el tipo de documento donde aparece.
# "parte0" y "parte1" son las partes detectadas; "emisor", quien firma o emite.
ROLES = {
    "DEMANDA":           {"parte0": "Demandante", "parte1": "Demandado",
                          "emisor": "Abogado del demandante"},
    "CONTESTACION":      {"parte0": "Demandado", "emisor": "Abogado del demandado"},
    "CITACION":          {"emisor": "Tribunal"},
    "ORDEN":             {"emisor": "Tribunal"},
    "MOCION":            {"emisor": "Abogado"},
    "ESCRITO":           {"emisor": "Abogado"},
    "NOTIFICACION":      {"emisor": "Abogado"},
    "ACUSE":             {"emisor": "Notificador"},
    "RECLAMACION":       {"parte0": "Reclamante", "emisor": "Agencia"},
    "SEGURO":            {"emisor": "Aseguradora"},
    "EXPEDIENTE-MEDICO": {"emisor": "Proveedor medico"},
    "FACTURA":           {"emisor": "Proveedor"},
    "REPORTE-POLICIAL":  {"emisor": "Autoridad"},
    "DECLARACION":       {"emisor": "Declarante"},
    "CONTRATO":          {"parte0": "Parte", "parte1": "Parte"},
    "CORRESPONDENCIA":   {"emisor": "Corresponsal"},
    "TRANSCRIPCION":     {"emisor": "Transcriptor"},
    "DESCUBRIMIENTO":    {"emisor": "Abogado"},
    "PRUEBA":            {},
    "CARATULA":          {},
}

RESULTADO_ATENDIDA = re.compile(r"connected|answered|atendida", re.IGNORECASE)

# Los nombres que salen de un documento llegan con restos del formato: el
# "RE:" de una carta, la coletilla del numero de caso, una etiqueta de campo.
PREFIJOS = re.compile(r"^\s*(re|ref|asunto|subject|owner|patient name|claimant|"
                      r"name|nombre|para|to|attn)\s*:\s*", re.IGNORECASE)
COLETILLA = re.compile(r"\s*[-–—]\s*(case|caso|expediente|civil action)\s*(no\.?|n[uú]m)?.*$",
                       re.IGNORECASE)


def limpiar_nombre(crudo: str) -> str:
    """Quita del nombre lo que es formato del documento, no la persona."""
    texto = (crudo or "").strip()
    for _ in range(2):                      # puede venir con dos prefijos
        texto = PREFIJOS.sub("", texto).strip()
    texto = COLETILLA.sub("", texto).strip(" .,;:-")
    texto = re.sub(r"\s{2,}", " ", texto)
    return texto if 2 < len(texto) < 90 else ""


def _registrar(fuente: str, detalle: str, creados: int, actualizados: int,
               interacciones: int) -> Dict[str, Any]:
    db.execute(
        """INSERT INTO importaciones (fuente, detalle, creados, actualizados,
                                      interacciones, ts)
           VALUES (?,?,?,?,?,?)""",
        (fuente, detalle, creados, actualizados, interacciones, db.now_iso()))
    return {"fuente": fuente, "detalle": detalle, "creados": creados,
            "actualizados": actualizados, "interacciones": interacciones}


# ==========================================================================
# 1. Nomina del despacho
# ==========================================================================
def ruta_nomina() -> Optional[Path]:
    """El Excel de nomina, se llame como se llame."""
    return next(iter(sorted(SAMPLE_DIR.glob("*Biometric*.xlsx"))), None)


def ruta_llamadas() -> Optional[Path]:
    """El Excel de la central telefonica."""
    return next(iter(sorted(SAMPLE_DIR.glob("*RingCentral*.xlsx"))), None)


def importar_nomina(ruta: Optional[Path] = None) -> Dict[str, Any]:
    ruta = ruta or ruta_nomina()
    if not ruta or not Path(ruta).exists():
        return _registrar("nomina", "no se encontro el Excel de nomina", 0, 0, 0)

    import openpyxl
    wb = openpyxl.load_workbook(str(ruta), read_only=True, data_only=True)
    if "Nomina" not in wb.sheetnames:
        wb.close()
        return _registrar("nomina", "el Excel no tiene hoja 'Nomina'", 0, 0, 0)

    creados = actualizados = 0
    filas = list(wb["Nomina"].iter_rows(values_only=True))
    for fila in filas[1:]:
        if not fila or not fila[0] or not fila[1]:
            continue
        nombre = str(fila[0]).strip()
        ext = str(fila[1]).split(".")[0].strip()
        area = str(fila[2] or "").strip()
        if not nombre or not ext.isdigit():
            continue
        persona_id, nueva = modelo.asegurar_persona(
            nombre, identificadores=[("extension", ext, "central telefonica")],
            organizacion="Despacho", rol=area, origen="nomina", interno=True)
        modelo.actualizar_persona(persona_id, interno=1)
        creados += nueva
        actualizados += (not nueva)
    wb.close()
    return _registrar("nomina", f"{Path(ruta).name}", creados, actualizados, 0)


# ==========================================================================
# 2. Central telefonica
# ==========================================================================
def importar_llamadas(ruta: Optional[Path] = None) -> Dict[str, Any]:
    ruta = ruta or ruta_llamadas()
    if not ruta or not Path(ruta).exists():
        return _registrar("llamadas", "no se encontro el Excel de llamadas", 0, 0, 0)

    import openpyxl
    wb = openpyxl.load_workbook(str(ruta), read_only=True, data_only=True)

    llamadas: List[Dict[str, Any]] = []
    for hoja in wb.worksheets:
        filas = hoja.iter_rows(values_only=True)
        cabecera = next(filas, None)
        if not cabecera:
            continue
        col = {str(c).strip().lower(): i for i, c in enumerate(cabecera) if c}

        def val(fila, clave):
            i = col.get(clave, -1)
            return str(fila[i]).strip() if 0 <= i < len(fila) and fila[i] is not None else ""

        for fila in filas:
            if not fila or not any(fila):
                continue
            crudo_fecha = val(fila, "date")
            m = re.search(r"(\d{2})/(\d{2})/(\d{4})", crudo_fecha)
            if not m:
                continue
            hora = val(fila, "time")
            hm = re.match(r"(\d{1,2}):(\d{2})(?::(\d{2}))?", hora)
            t = dtime(int(hm.group(1)), int(hm.group(2)), int(hm.group(3) or 0)) if hm else dtime()
            llamadas.append({
                "ts": datetime.combine(date(int(m.group(3)), int(m.group(1)), int(m.group(2))), t),
                "desde": val(fila, "from"),
                "hacia": val(fila, "to"),
                "direccion": val(fila, "direction"),
                "extension": val(fila, "extension"),
                "resultado": val(fila, "action result"),
                "duracion": _segundos(val(fila, "duration")),
                "hoja": hoja.title,
            })
    wb.close()
    if not llamadas:
        return _registrar("llamadas", "el Excel no tenia llamadas legibles", 0, 0, 0)

    # --- decidir que numeros merecen una ficha --------------------------
    min_total = int(CONFIG.get("llamadas_min_segundos_total") or 900)
    min_veces = int(CONFIG.get("llamadas_min_veces") or 2)
    por_numero: Dict[str, List[Dict[str, Any]]] = defaultdict(list)
    for ll in llamadas:
        externo = ll["desde"] if ll["direccion"].lower().startswith("in") else ll["hacia"]
        norma = modelo.norm_telefono(externo)
        if len(norma) == 10:
            ll["externo"] = externo
            por_numero[norma].append(ll)

    creados = actualizados = hitos = 0
    for norma, grupo in por_numero.items():
        atendidas = [g for g in grupo if RESULTADO_ATENDIDA.search(g["resultado"])]
        conversado = sum(g["duracion"] for g in atendidas)
        # Entra al directorio quien parece una relacion, no cualquiera que marco:
        # llamo mas de una vez, o acumulo una conversacion larga.
        relevante = len(grupo) >= min_veces or conversado >= min_total
        if not (relevante and CONFIG.get("crear_desde_llamadas")):
            # No se crea ficha, pero el historial no se pierde: queda indexado
            # por numero para que una llamada entrante encuentre su contexto.
            for ll in sorted(grupo, key=lambda g: g["ts"], reverse=True)[:10]:
                ref = f"{norma}|{ll['ts'].isoformat()}"
                if modelo.registrar_interaccion(
                        None, "llamada", ll["ts"].isoformat(timespec="seconds"),
                        f"Llamada {ll['direccion'].lower()} de "
                        f"{modelo.formatear_telefono(ll['externo'])}",
                        detalle=f"{ll['extension'] or 'sin extension'} · {ll['resultado']}",
                        origen="RingCentral", referencia=ref,
                        meta={"numero": norma, "duracion": ll["duracion"],
                              "sin_ficha": True}):
                    hitos += 1
            continue

        muestra = grupo[0]
        persona_id, nueva = modelo.asegurar_persona(
            modelo.formatear_telefono(muestra["externo"]),
            identificadores=[("telefono", muestra["externo"], "llamadas")],
            rol="Contacto telefonico", origen="llamadas")
        creados += nueva
        actualizados += (not nueva)

        for ll in sorted(grupo, key=lambda g: g["ts"], reverse=True)[:25]:
            atendida = bool(RESULTADO_ATENDIDA.search(ll["resultado"]))
            quien = ll["extension"] or "sin extension"
            titulo = (f"Llamada {ll['direccion'].lower()} · "
                      f"{_duracion_legible(ll['duracion'])}" if atendida
                      else f"Llamada {ll['direccion'].lower()} sin atender")
            detalle = f"{quien} · {ll['resultado']}"
            ref = f"{norma}|{ll['ts'].isoformat()}"
            if modelo.registrar_interaccion(
                    persona_id, "llamada", ll["ts"].isoformat(timespec="seconds"),
                    titulo, detalle=detalle, origen="RingCentral", referencia=ref,
                    meta={"duracion": ll["duracion"], "resultado": ll["resultado"],
                          "extension": quien}):
                hitos += 1

        # la llamada tambien cuenta en la ficha de quien la atendio
        for ll in grupo:
            ext = (ll["extension"] or "").split("-")[0].strip()
            if not ext.isdigit():
                continue
            interno = modelo.buscar_por_identificador("extension", ext)
            if not interno:
                continue
            ref = f"ext{ext}|{norma}|{ll['ts'].isoformat()}"
            if modelo.registrar_interaccion(
                    interno["id"], "llamada", ll["ts"].isoformat(timespec="seconds"),
                    f"Atendio a {modelo.formatear_telefono(ll['externo'])}"
                    if RESULTADO_ATENDIDA.search(ll["resultado"])
                    else f"Llamada perdida de {modelo.formatear_telefono(ll['externo'])}",
                    detalle=ll["resultado"], origen="RingCentral", referencia=ref):
                hitos += 1

    return _registrar("llamadas", f"{Path(ruta).name} · {len(llamadas)} llamadas leidas",
                      creados, actualizados, hitos)


def _segundos(texto: str) -> int:
    partes = (texto or "").split(":")
    try:
        numeros = [int(p) for p in partes]
    except ValueError:
        return 0
    while len(numeros) < 3:
        numeros.insert(0, 0)
    return numeros[0] * 3600 + numeros[1] * 60 + numeros[2]


def _duracion_legible(segundos: int) -> str:
    if segundos < 60:
        return f"{segundos} s"
    return f"{segundos // 60} min"


# ==========================================================================
# 3. Documentos clasificados (base de datos de DocuFlow AI)
# ==========================================================================
def importar_documentos(ruta_db: Optional[Path] = None) -> Dict[str, Any]:
    ruta_db = Path(ruta_db or DOCS_DB)
    if not ruta_db.exists():
        return _registrar("documentos",
                          "no se encontro la base de DocuFlow AI "
                          "(data/documentos/documentos.db)", 0, 0, 0)

    conn = sqlite3.connect(f"file:{ruta_db}?mode=ro", uri=True)
    conn.row_factory = sqlite3.Row
    try:
        filas = [dict(r) for r in conn.execute(
            """SELECT id, nombre_final, nombre_original, categoria, categoria_nombre,
                      expediente, fecha_doc, titulo, entidades_json, estado
               FROM documentos WHERE estado <> 'duplicado' ORDER BY fecha_doc""")]
    except sqlite3.Error as exc:
        conn.close()
        return _registrar("documentos", f"no se pudo leer: {exc}", 0, 0, 0)
    conn.close()

    creados = actualizados = hitos = 0
    for doc in filas:
        try:
            ent = json.loads(doc["entidades_json"] or "{}")
        except json.JSONDecodeError:
            ent = {}
        categoria = doc["categoria"] or ""
        papeles = ROLES.get(categoria, {})
        caso_id = 0
        if doc["expediente"]:
            partes = [limpiar_nombre(p) for p in (ent.get("partes") or [])]
            partes = [p for p in partes if p]
            # el nombre del caso se toma de la demanda; los demas escritos no lo pisan
            nombre_caso = (" v. ".join(partes[:2])
                           if len(partes) >= 2 and categoria in ("DEMANDA", "CONTESTACION")
                           else "")
            caso_id = modelo.asegurar_caso(doc["expediente"], nombre=nombre_caso,
                                           tribunal=ent.get("tribunal", ""),
                                           monto=ent.get("monto_principal") or None)

        implicados: List[Tuple[str, str, str, List[Tuple[str, str, str]]]] = []
        for indice, parte in enumerate((ent.get("partes") or [])[:2]):
            rol = papeles.get(f"parte{indice}", "")
            limpio = limpiar_nombre(parte)
            if limpio:
                implicados.append((limpio, rol, "", []))
        if ent.get("emisor"):
            ids: List[Tuple[str, str, str]] = []
            for correo in (ent.get("correos") or [])[:1]:
                ids.append(("correo", correo, "documento"))
            for telefono in (ent.get("telefonos") or [])[:2]:
                ids.append(("telefono", telefono, "documento"))
            emisor = limpiar_nombre(ent["emisor"])
            if emisor:
                implicados.append((emisor, papeles.get("emisor", ""), emisor, ids))

        for nombre, rol, organizacion, ids in implicados:
            if not (nombre or "").strip():
                continue
            persona_id, nueva = modelo.asegurar_persona(
                nombre, identificadores=ids, organizacion=organizacion, rol=rol,
                origen="documentos")
            creados += nueva
            actualizados += (not nueva)
            if caso_id:
                modelo.asegurar_participacion(
                    persona_id, caso_id, rol, origen=f"documento {categoria}",
                    confianza=0.9 if rol else 0.4)
            titulo = doc["titulo"] or doc["nombre_original"]
            if modelo.registrar_interaccion(
                    persona_id, "documento",
                    (doc["fecha_doc"] or db.now_iso()[:10]) + "T12:00:00",
                    f"{doc['categoria_nombre'] or categoria}: {titulo}",
                    detalle=doc["nombre_final"] or doc["nombre_original"],
                    caso_id=caso_id or None, origen="DocuFlow AI",
                    referencia=f"doc{doc['id']}",
                    meta={"categoria": categoria}):
                hitos += 1

    _actualizar_conteo_casos()
    return _registrar("documentos", f"{len(filas)} documentos leidos",
                      creados, actualizados, hitos)


def _actualizar_conteo_casos() -> None:
    for caso in db.query("SELECT id FROM casos"):
        fila = db.query_one(
            """SELECT COUNT(*) n, MIN(ts) desde, MAX(ts) hasta FROM interacciones
               WHERE caso_id=? AND tipo='documento'""", (caso["id"],))
        db.execute("UPDATE casos SET documentos=?, desde=?, hasta=? WHERE id=?",
                   ((fila or {}).get("n", 0), (fila or {}).get("desde"),
                    (fila or {}).get("hasta"), caso["id"]))


# ==========================================================================
def importar_todo() -> List[Dict[str, Any]]:
    """Orden importante: la nomina primero, para poder ligar las extensiones."""
    return [importar_nomina(), importar_documentos(), importar_llamadas()]
