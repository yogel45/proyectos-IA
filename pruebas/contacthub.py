"""Todo lo que hay que saber para medir ContactHub: donde esta, como se entra
y con que datos se llena antes de empezar.

ContactHub no es una aplicacion abierta como las de los retos anteriores: cada
peticion (salvo `/health` y `/api/v1/auth/*`) exige una cabecera
`Authorization: Bearer ...`, y cada contacto pertenece a un usuario. Eso cambia
la prueba de carga en tres cosas:

* hay que abrir sesion antes de medir, y renovar el token porque caduca a los
  15 minutos — una sesion de pruebas dura mas que eso;
* la libreta empieza vacia, asi que hay que sembrarla: medir busquedas sobre
  una base sin contactos no mide nada;
* la siembra se hace por la propia importacion de ContactHub (CSV de Google
  Contacts), no escribiendo en su base de datos por detras. Asi la prueba usa
  el sistema por donde se usa de verdad.

Los datos de la siembra salen de los mismos archivos reales de la oficina que
alimentan el resto del proyecto (la nomina y el historial de la centralita),
leidos aqui directamente con openpyxl: este paquete no importa codigo de las
otras aplicaciones.
"""
from __future__ import annotations

import csv
import io
import os
import re
import subprocess
import sys
import time
import unicodedata
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import httpx

RAIZ = Path(__file__).resolve().parent.parent
SAMPLE = RAIZ / "sample_data"

PUERTO = int(os.environ.get("CH_PUERTO", "8765"))
BASE = os.environ.get("CH_BASE", f"http://127.0.0.1:{PUERTO}").rstrip("/")
API = f"{BASE}/api/v1"

# Cuenta que usan las pruebas. Es una cuenta mas, creada por el endpoint
# publico de registro: la prueba no toca la base de datos por detras.
USUARIO = os.environ.get("CH_USUARIO", "pruebas.carga@grupoilalo.com")
CLAVE = os.environ.get("CH_CLAVE", "CargaContactHub2026!")


# --------------------------------------------------------------------------
# Donde esta ContactHub
# --------------------------------------------------------------------------
CANDIDATAS = [
    "ContactHub", "../ContactHub", "../ContactHub/ContactHub",
    "../ContactHub-entrega/ContactHub", "ContactHub-entrega/ContactHub",
    "~/ContactHub", "~/Desktop/ContactHub", "~/Escritorio/ContactHub",
    "~/Downloads/ContactHub", "~/Descargas/ContactHub",
]


def localizar(pista: Optional[str] = None) -> Optional[Path]:
    """Busca la carpeta de ContactHub. Devuelve None si no aparece."""
    intentos = [pista] if pista else []
    intentos.append(os.environ.get("CONTACTHUB_DIR"))
    intentos += CANDIDATAS
    for bruto in intentos:
        if not bruto:
            continue
        ruta = Path(bruto).expanduser()
        if not ruta.is_absolute():
            ruta = (RAIZ / ruta).resolve()
        if (ruta / "app" / "main.py").exists():
            return ruta
    return None


def responde(timeout: float = 1.5) -> bool:
    try:
        return httpx.get(f"{BASE}/health", timeout=timeout).status_code == 200
    except Exception:
        return False


def arrancar(log: Path, pista: Optional[str] = None) -> Tuple[Optional[subprocess.Popen], str]:
    """Levanta ContactHub si no estaba levantado.

    Devuelve (proceso, motivo). El proceso es None cuando ya corria — en ese
    caso no se apaga al terminar: no es nuestro.
    """
    if responde():
        return None, "ya estaba corriendo"
    carpeta = localizar(pista)
    if carpeta is None:
        return None, "no se encontro la carpeta de ContactHub"
    if not (carpeta / ".env").exists():
        # ContactHub trae un script que genera las claves; sin .env no arranca.
        subprocess.run([sys.executable, "-m", "scripts.init_env"], cwd=str(carpeta),
                       capture_output=True, timeout=120)
    log.parent.mkdir(parents=True, exist_ok=True)
    salida = open(log, "w")
    proc = subprocess.Popen(
        [sys.executable, "-m", "uvicorn", "app.main:app_factory", "--factory",
         "--host", "127.0.0.1", "--port", str(PUERTO), "--log-level", "warning"],
        cwd=str(carpeta), stdout=salida, stderr=subprocess.STDOUT,
        start_new_session=(os.name != "nt"))
    return proc, f"arrancado desde {carpeta}"


def esperar(segundos: float = 60) -> float:
    """Espera a que /health conteste. Devuelve lo que tardo, o -1."""
    t0 = time.time()
    limite = t0 + segundos
    while time.time() < limite:
        if responde(timeout=2):
            return round(time.time() - t0, 2)
        time.sleep(0.25)
    return -1.0


# --------------------------------------------------------------------------
# Sesion: el token y su renovacion
# --------------------------------------------------------------------------
class Sesion:
    """Guarda el token de acceso y lo renueva antes de que caduque.

    ContactHub da tokens de 15 minutos. Una sesion de pruebas completa dura
    mas, y un 401 a mitad de la escalada se contaria como error del sistema
    cuando en realidad seria un error de la prueba.
    """

    def __init__(self, usuario: str = USUARIO, clave: str = CLAVE):
        self.usuario = usuario
        self.clave = clave
        self.token = ""
        self.refresh = ""
        self.caduca = 0.0
        self.usuario_id = ""
        self.renovaciones = 0

    def abrir(self) -> Dict[str, Any]:
        """Registra la cuenta si hace falta y abre sesion."""
        with httpx.Client(timeout=30) as c:
            r = c.post(f"{API}/auth/register", json={
                "email": self.usuario, "password": self.clave,
                "full_name": "Pruebas de carga"})
            nueva = r.status_code == 201
            if not nueva and r.status_code not in (409, 400, 403):
                return {"ok": False, "motivo": f"registro HTTP {r.status_code}: {r.text[:180]}"}
            r = c.post(f"{API}/auth/login",
                       json={"email": self.usuario, "password": self.clave})
            if r.status_code != 200:
                return {"ok": False,
                        "motivo": f"login HTTP {r.status_code}: {r.text[:180]}"}
            self._guardar(r.json())
            yo = c.get(f"{API}/auth/me", headers=self.cabeceras())
            if yo.status_code == 200:
                self.usuario_id = yo.json().get("id", "")
        return {"ok": True, "cuenta_nueva": nueva, "usuario": self.usuario,
                "usuario_id": self.usuario_id,
                "vida_token_s": int(self.caduca - time.time())}

    def _guardar(self, datos: Dict[str, Any]) -> None:
        self.token = datos["access_token"]
        self.refresh = datos.get("refresh_token", "")
        self.caduca = time.time() + float(datos.get("expires_in") or 900)

    def cabeceras(self) -> Dict[str, str]:
        """El Bearer, renovado si le queda poca vida."""
        if self.token and time.time() > self.caduca - 120:
            self.renovar()
        return {"Authorization": f"Bearer {self.token}"}

    def renovar(self) -> bool:
        with httpx.Client(timeout=30) as c:
            if self.refresh:
                r = c.post(f"{API}/auth/refresh", json={"refresh_token": self.refresh})
                if r.status_code == 200:
                    self._guardar(r.json())
                    self.renovaciones += 1
                    return True
            r = c.post(f"{API}/auth/login",
                       json={"email": self.usuario, "password": self.clave})
            if r.status_code == 200:
                self._guardar(r.json())
                self.renovaciones += 1
                return True
        return False


# --------------------------------------------------------------------------
# Siembra: una libreta de verdad, sacada de los archivos de la oficina
# --------------------------------------------------------------------------
_CIUDADES = [("Charleston", "SC"), ("Columbia", "SC"), ("Greenville", "SC"),
             ("Charlotte", "NC"), ("Atlanta", "GA"), ("Savannah", "GA"),
             ("Miami", "FL"), ("Orlando", "FL")]
_CATEGORIAS = ["Cliente", "Contraparte", "Juzgado", "Proveedor", "Interno"]
_ETIQUETAS = ["Nomina", "Centralita", "Urgente", "Seguimiento", "Archivo"]


def _sin_acentos(texto: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFD", texto)
                   if unicodedata.category(c) != "Mn")


def _correo(nombre: str, apellido: str, i: int, dominio: str) -> str:
    n = re.sub(r"[^a-z]", "", _sin_acentos(nombre).lower()) or "contacto"
    a = re.sub(r"[^a-z]", "", _sin_acentos(apellido).lower()) or "oficina"
    return f"{n}.{a}{i}@{dominio}"


def _leer_nomina() -> List[Tuple[str, str, str]]:
    """Los empleados reales: (nombre, apellido, area).

    Se lee la hoja "Nomina" del Excel del biometrico, que es una tabla
    limpia: columna "Names" y columna "Area". No se rastrea el libro entero
    buscando algo que parezca un nombre — esa heuristica colaba encabezados
    como "Hora Entrada" y ensuciaba los datos de la prueba.
    """
    ruta = next(iter(sorted(SAMPLE.glob("*Biometric*.xlsx"))), None)
    if ruta is None:
        return []
    try:
        from openpyxl import load_workbook
    except ImportError:
        return []
    salida: List[Tuple[str, str, str]] = []
    try:
        libro = load_workbook(ruta, read_only=True, data_only=True)
        hoja = libro["Nomina"] if "Nomina" in libro.sheetnames else libro.worksheets[0]
        filas = list(hoja.iter_rows(values_only=True))
        libro.close()
    except Exception:
        return []
    if not filas:
        return []
    encabezado = [str(c or "").strip().lower() for c in filas[0]]
    col_nombre = encabezado.index("names") if "names" in encabezado else 0
    col_area = encabezado.index("area") if "area" in encabezado else None
    for fila in filas[1:]:
        bruto = str(fila[col_nombre] or "").strip()
        if not bruto or any(ch.isdigit() for ch in bruto):
            continue
        partes = bruto.split()
        if len(partes) < 2:
            continue
        area = str(fila[col_area] or "").strip() if col_area is not None else ""
        # Debajo de la tabla hay una nota al pie ("Hora Sabado", "8am - 14 pm")
        # escrita en la columna de nombres. La segunda linea se cae sola por
        # llevar digitos; la primera no, y durante varias sesiones "Hora
        # Sabado" fue una empleada mas. Una fila de nomina de verdad tiene
        # area: la nota no.
        if not area:
            continue
        salida.append((partes[0].title(), " ".join(partes[1:]).title(), area))
    return salida


def _leer_telefonos() -> List[str]:
    """Numeros distintos del historial de la centralita."""
    ruta = next(iter(sorted(SAMPLE.glob("*RingCentral*.xlsx"))), None)
    if ruta is None:
        return []
    try:
        from openpyxl import load_workbook
    except ImportError:
        return []
    numeros: List[str] = []
    vistos = set()
    try:
        libro = load_workbook(ruta, read_only=True, data_only=True)
        for hoja in libro.worksheets:
            for fila in hoja.iter_rows(values_only=True):
                for celda in fila:
                    texto = str(celda or "")
                    digitos = re.sub(r"\D", "", texto)
                    if len(digitos) == 11 and digitos.startswith("1"):
                        digitos = digitos[1:]
                    if len(digitos) != 10 or digitos in vistos:
                        continue
                    vistos.add(digitos)
                    numeros.append(f"+1{digitos}")
                    if len(numeros) >= 4000:
                        libro.close()
                        return numeros
        libro.close()
    except Exception:
        return numeros
    return numeros


_RESERVA_NOMBRES = [
    ("Maria", "Gonzalez", "Paralegal"), ("Jose", "Rodriguez", "Marketing"),
    ("Ana", "Martinez", "Paralegal"), ("Luis", "Hernandez", "Data"),
    ("Carmen", "Lopez", "Paralegal"), ("Miguel", "Perez", "Data"),
    ("Rosa", "Sanchez", "Contabilidad"), ("Juan", "Ramirez", "Marketing"),
    ("Laura", "Torres", "Paralegal"), ("Pedro", "Flores", "Recepcion"),
    ("Sofia", "Rivera", "Paralegal"), ("Diego", "Gomez", "Data"),
    ("Elena", "Diaz", "Contabilidad"), ("Carlos", "Cruz", "Marketing"),
    ("Patricia", "Morales", "Paralegal"), ("Andres", "Ortiz", "Data"),
    ("Lucia", "Castillo", "Paralegal"), ("Javier", "Vargas", "Marketing"),
    ("Isabel", "Romero", "Recepcion"), ("Raul", "Mendoza", "Data"),
    ("Silvia", "Guerrero", "Paralegal"), ("Tomas", "Navarro", "Contabilidad"),
    ("Beatriz", "Campos", "Paralegal"), ("Ruben", "Aguilar", "Marketing"),
    ("Natalia", "Fuentes", "Recepcion"),
]
_EMPRESAS = [
    "Grupo Ilalo", "Coastal Freight LLC", "Palmetto Staffing",
    "Lowcountry Roofing", "Atlantic Seafood Co", "Southern Drywall Inc",
    "Carolina Landscaping", "Harbor Logistics", "Magnolia Hospitality",
    "Blue Ridge Contractors", "Tribunal Federal", "Oficina del Fiscal",
]

COLUMNAS_CSV = ["First Name", "Last Name", "Nickname", "Organization Name",
                "Organization Title", "Birthday", "Notes", "Labels",
                "E-mail 1 - Label", "E-mail 1 - Value",
                "Phone 1 - Label", "Phone 1 - Value",
                "Phone 2 - Label", "Phone 2 - Value",
                "Address 1 - Label", "Address 1 - Street", "Address 1 - City",
                "Address 1 - Region", "Address 1 - Postal Code",
                "Address 1 - Country",
                "Custom Field 1 - Label", "Custom Field 1 - Value"]


def csv_de_siembra(cuantos: int, desde: int = 0) -> bytes:
    """Construye un CSV en formato Google Contacts con `cuantos` contactos.

    Los nombres, los puestos y los telefonos salen de los archivos reales de
    la oficina; si no estan, se usa una lista de reserva para que la prueba
    siga siendo ejecutable en cualquier maquina.

    Los apellidos se combinan de dos en dos (apellido paterno y materno, como
    en los nombres hispanos reales de la nomina). Eso da cientos de nombres
    distintos sin inventar palabras: las busquedas de la prueba usan asi
    apellidos de verdad, y no cadenas como "Contacto 1423" que cualquier
    indice resuelve sin esfuerzo.
    """
    nombres = _leer_nomina() or list(_RESERVA_NOMBRES)
    telefonos = _leer_telefonos()

    buffer = io.StringIO(newline="")
    w = csv.DictWriter(buffer, fieldnames=COLUMNAS_CSV)
    w.writeheader()

    for n in range(cuantos):
        i = desde + n
        # El segundo apellido tiene que avanzar a distinto ritmo que el
        # primero. Un salto fijo no sirve: con 26 nombres, (i*13) % 26 solo
        # toma dos valores, y la primera version de esto genero 2 000
        # contactos con dos apellidos maternos entre todos. Usando la vuelta
        # de la lista salen len(nombres)^2 nombres completos distintos.
        vuelta = i // len(nombres)
        nombre, paterno, area = nombres[i % len(nombres)]
        materno = nombres[(i + vuelta) % len(nombres)][1]
        empresa = _EMPRESAS[vuelta % len(_EMPRESAS)]
        ciudad, region = _CIUDADES[i % len(_CIUDADES)]
        dominio = re.sub(r"[^a-z]", "", _sin_acentos(empresa).lower())[:18] + ".com"
        tel1 = telefonos[i % len(telefonos)] if telefonos else f"+1843555{i % 10000:04d}"
        tel2 = telefonos[(i * 7 + 3) % len(telefonos)] if telefonos else ""
        etiquetas = " ::: ".join(
            [_ETIQUETAS[i % len(_ETIQUETAS)]] + (["Starred"] if i % 17 == 0 else []))
        w.writerow({
            "First Name": nombre,
            "Last Name": f"{paterno} {materno}",
            "Nickname": _sin_acentos(nombre)[:3].lower() + str(i % 1000),
            "Organization Name": empresa,
            "Organization Title": area,
            "Birthday": f"19{60 + i % 40:02d}-{1 + i % 12:02d}-{1 + i % 28:02d}",
            "Notes": f"Contacto de la libreta de la oficina. Registro {i}.",
            "Labels": etiquetas,
            "E-mail 1 - Label": "Work",
            "E-mail 1 - Value": _correo(nombre, paterno, i, dominio),
            "Phone 1 - Label": "Work",
            "Phone 1 - Value": tel1,
            "Phone 2 - Label": "Mobile",
            "Phone 2 - Value": tel2,
            "Address 1 - Label": "Work",
            "Address 1 - Street": f"{100 + i % 900} Main St",
            "Address 1 - City": ciudad,
            "Address 1 - Region": region,
            "Address 1 - Postal Code": f"{29400 + i % 500}",
            "Address 1 - Country": "United States",
            "Custom Field 1 - Label": "Category",
            "Custom Field 1 - Value": _CATEGORIAS[i % len(_CATEGORIAS)],
        })
    return buffer.getvalue().encode("utf-8")


def sembrar(sesion: Sesion, objetivo: int = 2000,
            espera_max: float = 300) -> Dict[str, Any]:
    """Deja la libreta con al menos `objetivo` contactos.

    Si ya los tiene, no hace nada: repetir la prueba no debe duplicar la base.
    La importacion se encola y corre en segundo plano, asi que se espera al
    trabajo y de paso se mide lo que tarda.
    """
    with httpx.Client(timeout=120) as c:
        stats = c.get(f"{API}/stats", headers=sesion.cabeceras())
        total = stats.json().get("total", 0) if stats.status_code == 200 else 0
        if total >= objetivo:
            return {"sembrado": False, "total": total,
                    "motivo": f"la libreta ya tenia {total} contactos"}

        faltan = objetivo - total
        datos = csv_de_siembra(faltan, desde=total)
        t0 = time.perf_counter()
        r = c.post(f"{API}/imports/google-csv",
                   headers=sesion.cabeceras(),
                   files={"file": ("libreta.csv", datos, "text/csv")},
                   data={"on_conflict": "fill_missing"})
        if r.status_code != 202:
            return {"sembrado": False, "total": total,
                    "motivo": f"la importacion respondio HTTP {r.status_code}: {r.text[:200]}"}
        job = r.json().get("id")
        estado: Dict[str, Any] = {}
        limite = time.time() + espera_max
        while time.time() < limite:
            j = c.get(f"{API}/imports/{job}", headers=sesion.cabeceras())
            if j.status_code == 200:
                estado = j.json()
                if estado.get("status") in ("done", "completed", "finished", "error", "failed"):
                    break
            time.sleep(1.0)
        segundos = round(time.perf_counter() - t0, 1)
        stats = c.get(f"{API}/stats", headers=sesion.cabeceras())
        final = stats.json().get("total", 0) if stats.status_code == 200 else 0
    return {"sembrado": True, "pedidos": faltan, "total": final,
            "antes": total, "segundos": segundos,
            "contactos_por_segundo": round((final - total) / max(segundos, 0.1), 1),
            "trabajo": {k: estado.get(k) for k in
                        ("status", "created", "updated", "skipped", "errors", "total")}}


def inventario(sesion, cuantos: int = 150) -> Dict[str, Any]:
    """Averigua contra que se va a medir: ids, textos de busqueda, etiquetas.

    Se leen del sistema en vez de darlos por supuestos. Y se leen *repartidos
    por toda la libreta*, no de la primera pagina: pedir los 150 primeros
    ordenados por nombre devuelve 150 personas que se apellidan casi igual, y
    entonces todas las busquedas de la prueba caerian sobre la misma fila. La
    primera version de esta funcion hacia justo eso, y de 2 000 contactos
    salian dos apellidos distintos.

    Si la libreta esta vacia, se dice en voz alta: los 404 que saldrian
    entonces serian fallos de la prueba, no del sistema.
    """
    ids: List[str] = []
    textos: List[str] = []
    empresas: List[str] = []
    etiquetas: List[str] = []
    total = 0
    with httpx.Client(timeout=60) as c:
        cab = sesion.cabeceras()
        r = c.get(f"{API}/stats", headers=cab)
        if r.status_code == 200:
            total = r.json().get("total", 0)

        paginas = 12
        por_pagina = max(10, cuantos // paginas)
        if total > por_pagina:
            saltos = [int(i * (total - por_pagina) / (paginas - 1))
                      for i in range(paginas)]
        else:
            saltos = [0]
        for salto in saltos:
            r = c.get(f"{API}/contacts?limit={por_pagina}&offset={salto}&sort=name",
                      headers=cab)
            if r.status_code != 200:
                continue
            for item in r.json().get("items", []):
                ids.append(item["id"])
                nombre = (item.get("display_name") or "").split()
                if nombre:
                    textos.append(nombre[-1])
                if item.get("company"):
                    empresas.append(item["company"])

        r = c.get(f"{API}/labels", headers=cab)
        if r.status_code == 200:
            etiquetas = [e["name"] for e in r.json()][:8]

    # se quitan repetidos conservando el orden
    ids = list(dict.fromkeys(ids))
    textos = list(dict.fromkeys(textos))
    empresas = list(dict.fromkeys(empresas))
    return {"total": total, "ids": ids, "textos": textos or ["a"],
            "empresas": empresas or ["Grupo"], "etiquetas": etiquetas,
            "con_datos": bool(ids and total >= 50)}


def limpiar(sesion: Sesion) -> Dict[str, int]:
    """Borra lo que escribieron los escenarios de escritura.

    Todo lo que crea la prueba lleva empresa "Prueba de carga": se localiza
    con el mismo filtro publico de la API y se borra en bloque. La libreta
    sembrada se queda: volver a sembrarla en cada corrida tarda mas que la
    propia prueba.
    """
    borrados = 0
    with httpx.Client(timeout=120) as c:
        for _ in range(60):                       # 60 x 200 = tope de 12 000
            cab = sesion.cabeceras()
            r = c.get(f"{API}/contacts?company=Prueba de carga&limit=200",
                      headers=cab)
            if r.status_code != 200:
                break
            ids = [x["id"] for x in r.json().get("items", [])]
            if not ids:
                break
            d = c.post(f"{API}/contacts/bulk", headers=cab,
                       json={"ids": ids, "action": "delete"})
            if d.status_code != 200:
                break
            borrados += d.json().get("affected", 0)
    return {"contactos_de_prueba_borrados": borrados}
