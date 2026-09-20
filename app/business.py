"""Integracion con datos operativos de la oficina.

El analisis de video no vive aislado: se cruza con dos fuentes reales del
negocio para que las metricas tengan sentido gerencial.

  * Biometrico (`Datos_Biometrico_2.xlsx`): nomina + marcajes diarios por mes.
  * RingCentral (`RingCentral_History.xlsx`): historico de llamadas por semana.

Con eso el dashboard responde preguntas como: "a las 10:00 habia pico de
llamadas, cuanta gente habia realmente en piso segun la camara y segun el
biometrico?".
"""
from __future__ import annotations

import logging
import re
from datetime import date, datetime, time as dtime, timedelta
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple

from . import db
from .config import SAMPLE_DIR

log = logging.getLogger("officevision.business")

MONTHS = {
    "january": 1, "february": 2, "febraury": 2, "march": 3, "april": 4, "may": 5,
    "june": 6, "july": 7, "august": 8, "september": 9, "october": 10,
    "november": 11, "december": 12,
}


def _as_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (dtime, datetime)):
        return value.strftime("%H:%M:%S")
    return str(value).strip()


def _parse_time(value: Any) -> Optional[dtime]:
    text = _as_text(value)
    if not text or text in {"-", "--"}:
        return None
    m = re.match(r"^(\d{1,2}):(\d{2})(?::(\d{2}))?$", text)
    if not m:
        return None
    h, mi, s = int(m.group(1)), int(m.group(2)), int(m.group(3) or 0)
    if h > 23 or mi > 59:
        return None
    return dtime(h, mi, s)


def _minutes(t: Optional[dtime]) -> Optional[float]:
    return None if t is None else t.hour * 60 + t.minute + t.second / 60.0


# ==========================================================================
# Importacion
# ==========================================================================
def import_biometric(path: Path) -> Dict[str, int]:
    """Carga la hoja `Nomina` y los marcajes de cada hoja mensual."""
    import openpyxl

    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    employees = 0
    records = 0

    if "Nomina" in wb.sheetnames:
        rows = list(wb["Nomina"].iter_rows(values_only=True))
        for row in rows[1:]:
            if not row or not row[0] or not row[1]:
                continue
            name = _as_text(row[0])
            emp_id = _as_text(row[1]).split(".")[0]
            if not name or not emp_id.isdigit():
                continue
            db.execute(
                """INSERT INTO employees (emp_id, name, area, entry_time, lunch, exit_time)
                   VALUES (?,?,?,?,?,?)
                   ON CONFLICT(emp_id) DO UPDATE SET
                     name=excluded.name, area=excluded.area,
                     entry_time=excluded.entry_time, lunch=excluded.lunch,
                     exit_time=excluded.exit_time""",
                (emp_id, name, _as_text(row[2]), _as_text(row[3]), _as_text(row[4]),
                 _as_text(row[5])),
            )
            employees += 1

    schedule = {e["emp_id"]: e for e in db.query("SELECT * FROM employees")}

    for sheet_name in wb.sheetnames:
        month = MONTHS.get(sheet_name.strip().lower())
        if not month:
            continue
        records += _import_month(wb[sheet_name], sheet_name, month, schedule)
    wb.close()
    return {"employees": employees, "attendance": records}


def _import_month(ws, sheet_name: str, month: int,
                  schedule: Dict[str, Dict[str, Any]]) -> int:
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    year = datetime.now().year
    day_cols: Dict[int, int] = {}
    header_row = -1

    for i, row in enumerate(rows[:12]):
        joined = " ".join(_as_text(c) for c in row[:4])
        m = re.search(r"From:\s*(\d{4})-(\d{2})-(\d{2})", joined)
        if m:
            year = int(m.group(1))
        if any(_as_text(c) == "No." for c in row[:3]):
            header_row = i
            for col, cell in enumerate(row):
                text = _as_text(cell)
                if text.isdigit() and 1 <= int(text) <= 31:
                    day_cols[int(text)] = col
    if header_row < 0 or not day_cols:
        return 0

    inserted = 0
    current: Optional[Dict[str, Any]] = None
    block: Dict[str, List[Any]] = {}

    def flush() -> int:
        nonlocal block, current
        if not current or "Check-in1" not in block:
            return 0
        n = 0
        emp_id = current["emp_id"]
        sched = schedule.get(emp_id, {})
        exp_in = _parse_time(sched.get("entry_time"))
        exp_out = _parse_time(sched.get("exit_time"))
        for day, col in sorted(day_cols.items()):
            try:
                d = date(year, month, day)
            except ValueError:
                continue
            cin = _parse_time(block.get("Check-in1", [None] * 60)[col]
                              if col < len(block.get("Check-in1", [])) else None)
            outs = block.get("Check-out1", [])
            cout = _parse_time(outs[col] if col < len(outs) else None)
            if cin is None and cout is None:
                continue
            worked = None
            if cin and cout:
                worked = max(0.0, (_minutes(cout) or 0) - (_minutes(cin) or 0))
            late = None
            if cin and exp_in:
                late = max(0.0, (_minutes(cin) or 0) - (_minutes(exp_in) or 0))
            early = None
            if cout and exp_out:
                early = max(0.0, (_minutes(exp_out) or 0) - (_minutes(cout) or 0))
            db.execute(
                """INSERT INTO attendance (emp_id, name, date, check_in, check_out,
                                           worked_min, late_min, early_min, month)
                   VALUES (?,?,?,?,?,?,?,?,?)
                   ON CONFLICT(emp_id, date) DO UPDATE SET
                     check_in=excluded.check_in, check_out=excluded.check_out,
                     worked_min=excluded.worked_min, late_min=excluded.late_min,
                     early_min=excluded.early_min""",
                (emp_id, current["name"], d.isoformat(),
                 cin.strftime("%H:%M:%S") if cin else None,
                 cout.strftime("%H:%M:%S") if cout else None,
                 worked, late, early, sheet_name),
            )
            n += 1
        block = {}
        current = None
        return n

    for row in rows[header_row + 1:]:
        if not row:
            continue
        label = _as_text(row[3]) if len(row) > 3 else ""
        no = _as_text(row[0]) if row else ""
        emp_id = _as_text(row[1]).split(".")[0] if len(row) > 1 else ""
        name = _as_text(row[2]) if len(row) > 2 else ""
        if label == "Check-in1" and emp_id:
            inserted += flush()
            current = {"emp_id": emp_id, "name": name or emp_id, "no": no}
            block = {}
        if current and label:
            block[label] = list(row)
    inserted += flush()
    return inserted


CALL_DATE_RE = re.compile(r"(\d{2})/(\d{2})/(\d{4})")


def _duration_seconds(value: Any) -> int:
    text = _as_text(value)
    if not text:
        return 0
    parts = text.split(":")
    try:
        parts = [int(p) for p in parts]
    except ValueError:
        return 0
    while len(parts) < 3:
        parts.insert(0, 0)
    return parts[0] * 3600 + parts[1] * 60 + parts[2]


def import_calls(path: Path) -> Dict[str, int]:
    """Carga los logs de RingCentral (una hoja por semana)."""
    import openpyxl

    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    total = 0
    db.execute("DELETE FROM calls")
    for sheet in wb.worksheets:
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            continue
        cols = {_as_text(c).lower(): i for i, c in enumerate(header) if c}
        batch = []
        for row in rows:
            if not row or not any(row):
                continue

            def val(key: str) -> str:
                i = cols.get(key, -1)
                return _as_text(row[i]) if 0 <= i < len(row) else ""

            raw_date = val("date")
            m = CALL_DATE_RE.search(raw_date)
            if not m:
                continue
            d = date(int(m.group(3)), int(m.group(1)), int(m.group(2)))
            t = _parse_time(val("time")) or dtime(0, 0)
            ext = val("extension")
            ext_id = ext.split("-")[0].strip() if "-" in ext else ""
            batch.append((
                datetime.combine(d, t).isoformat(timespec="seconds"), d.isoformat(),
                t.hour, val("direction"), val("from"), val("to"), ext, ext_id,
                val("name"), val("action result"),
                _duration_seconds(val("duration")), sheet.title,
            ))
        if batch:
            db.executemany(
                """INSERT INTO calls (ts, date, hour, direction, from_num, to_num,
                                      extension, ext_id, agent, result, duration_s, week)
                   VALUES (?,?,?,?,?,?,?,?,?,?,?,?)""", batch)
            total += len(batch)
    wb.close()
    return {"calls": total}


def import_samples(force: bool = False) -> Dict[str, Any]:
    """Importa los archivos de `sample_data/` si la BD esta vacia."""
    out: Dict[str, Any] = {}
    have_att = (db.query_one("SELECT COUNT(*) n FROM attendance") or {}).get("n", 0)
    have_calls = (db.query_one("SELECT COUNT(*) n FROM calls") or {}).get("n", 0)
    bio = next(iter(sorted(SAMPLE_DIR.glob("*Biometric*.xlsx"))), None)
    calls = next(iter(sorted(SAMPLE_DIR.glob("*RingCentral*.xlsx"))), None)
    if bio and (force or not have_att):
        out.update(import_biometric(bio))
    if calls and (force or not have_calls):
        out.update(import_calls(calls))
    return out


# ==========================================================================
# Analisis cruzado
# ==========================================================================
def calls_by_hour(day: Optional[str] = None) -> List[Dict[str, Any]]:
    where, params = "", []
    if day:
        where = "WHERE date = ?"
        params = [day]
    rows = db.query(
        f"""SELECT hour,
                   COUNT(*) AS llamadas,
                   SUM(CASE WHEN result LIKE '%connected%' THEN 1 ELSE 0 END) AS atendidas,
                   SUM(CASE WHEN result IN ('Missed','No answer','Rejected','Voicemail','Busy')
                            THEN 1 ELSE 0 END) AS perdidas,
                   ROUND(AVG(duration_s),1) AS duracion_prom
            FROM calls {where} GROUP BY hour ORDER BY hour""", params)
    return rows


def attendance_by_hour(day: Optional[str] = None) -> List[Dict[str, Any]]:
    """Personas presentes segun el biometrico, hora por hora."""
    if not day:
        row = db.query_one("SELECT MAX(date) AS d FROM attendance")
        day = (row or {}).get("d")
    if not day:
        return []
    rows = db.query(
        "SELECT check_in, check_out FROM attendance WHERE date=? AND check_in IS NOT NULL",
        (day,))
    buckets = {h: 0 for h in range(24)}
    for r in rows:
        cin = _parse_time(r["check_in"])
        cout = _parse_time(r["check_out"]) or dtime(23, 59)
        if not cin:
            continue
        for h in range(cin.hour, min(23, cout.hour) + 1):
            buckets[h] += 1
    return [{"hour": h, "presentes": n} for h, n in buckets.items()]


def detected_by_hour(day: Optional[str] = None) -> List[Dict[str, Any]]:
    """Personas detectadas por vision, promedio y pico por hora."""
    where, params = "", []
    if day:
        where = "WHERE substr(ts,1,10) = ?"
        params = [day]
    rows = db.query(
        f"""SELECT CAST(substr(ts,12,2) AS INTEGER) AS hour,
                   ROUND(AVG(persons),2) AS promedio,
                   MAX(persons) AS pico,
                   COUNT(*) AS muestras
            FROM snapshots {where} GROUP BY hour ORDER BY hour""", params)
    return rows


def schedule_expected_by_hour() -> List[Dict[str, Any]]:
    """Personal que *deberia* estar segun la nomina."""
    rows = db.query("SELECT entry_time, exit_time FROM employees")
    buckets = {h: 0 for h in range(24)}
    for r in rows:
        cin = _parse_time(r["entry_time"])
        cout = _parse_time(r["exit_time"])
        if not cin or not cout:
            continue
        for h in range(cin.hour, min(23, cout.hour) + 1):
            buckets[h] += 1
    return [{"hour": h, "esperados": n} for h, n in buckets.items()]


def common_days(limit: int = 60) -> List[str]:
    """Dias con datos de asistencia y de llamadas a la vez."""
    rows = db.query(
        """SELECT DISTINCT a.date AS d FROM attendance a
           JOIN calls c ON c.date = a.date ORDER BY d DESC LIMIT ?""", (limit,))
    return [r["d"] for r in rows]


def default_day() -> Optional[str]:
    days = common_days(1)
    if days:
        return days[0]
    row = db.query_one("SELECT MAX(date) AS d FROM attendance")
    return (row or {}).get("d")


def average_profile() -> Dict[int, Dict[str, float]]:
    """Perfil horario promedio (dia tipico) de llamadas y asistencia."""
    call_rows = db.query(
        """SELECT hour, COUNT(*) n,
                  SUM(CASE WHEN result LIKE '%connected%' THEN 1 ELSE 0 END) atendidas,
                  SUM(CASE WHEN result IN ('Missed','No answer','Rejected','Voicemail','Busy')
                           THEN 1 ELSE 0 END) perdidas
           FROM calls GROUP BY hour""")
    call_days = (db.query_one("SELECT COUNT(DISTINCT date) n FROM calls") or {}).get("n", 1) or 1
    att_rows = db.query(
        "SELECT date, check_in, check_out FROM attendance WHERE check_in IS NOT NULL")
    att_days: Dict[str, Dict[int, int]] = {}
    for r in att_rows:
        cin = _parse_time(r["check_in"])
        cout = _parse_time(r["check_out"]) or dtime(23, 59)
        if not cin:
            continue
        bucket = att_days.setdefault(r["date"], {})
        for h in range(cin.hour, min(23, cout.hour) + 1):
            bucket[h] = bucket.get(h, 0) + 1
    n_days = max(1, len(att_days))
    profile: Dict[int, Dict[str, float]] = {}
    for h in range(24):
        c = next((r for r in call_rows if r["hour"] == h), {})
        present = sum(d.get(h, 0) for d in att_days.values()) / n_days
        profile[h] = {
            "llamadas": round((c.get("n") or 0) / call_days, 1),
            "atendidas": round((c.get("atendidas") or 0) / call_days, 1),
            "perdidas": round((c.get("perdidas") or 0) / call_days, 1),
            "biometrico": round(present, 1),
        }
    return profile


def cross_analysis(day: Optional[str] = None, mode: str = "auto") -> Dict[str, Any]:
    """Une nomina + biometrico + llamadas + vision en una sola tabla horaria.

    mode='day'  -> un dia concreto; mode='avg' -> dia tipico (promedios).
    mode='auto' -> dia concreto si hay uno con ambas fuentes, si no promedio.
    """
    if mode == "auto":
        day = day or default_day()
        mode = "day" if day and common_days(90) and day in common_days(90) else "avg"
    expected = {r["hour"]: r["esperados"] for r in schedule_expected_by_hour()}
    detected = {r["hour"]: r for r in detected_by_hour()}

    if mode == "day" and day:
        present = {r["hour"]: r["presentes"] for r in attendance_by_hour(day)}
        calls = {r["hour"]: r for r in calls_by_hour(day)}
        source = {h: {"llamadas": (calls.get(h, {}).get("llamadas") or 0),
                      "atendidas": (calls.get(h, {}).get("atendidas") or 0),
                      "perdidas": (calls.get(h, {}).get("perdidas") or 0),
                      "biometrico": present.get(h, 0)} for h in range(24)}
        etiqueta = f"Dia {day}"
    else:
        source = average_profile()
        etiqueta = "Dia tipico (promedio historico)"

    table, findings = [], []
    for h in range(6, 21):
        s_h = source.get(h, {})
        d = detected.get(h, {})
        row = {
            "hour": f"{h:02d}:00",
            "esperados": expected.get(h, 0),
            "biometrico": s_h.get("biometrico", 0),
            "detectados": d.get("promedio") or 0,
            "pico_detectado": d.get("pico") or 0,
            "llamadas": s_h.get("llamadas", 0),
            "atendidas": s_h.get("atendidas", 0),
            "perdidas": s_h.get("perdidas", 0),
        }
        table.append(row)
        if row["llamadas"] >= 5 and row["esperados"] and row["biometrico"] < row["esperados"] * 0.6:
            findings.append(
                f"{row['hour']}: {row['llamadas']:.0f} llamadas con {row['biometrico']:.1f} "
                f"de {row['esperados']} personas presentes (cobertura baja).")
        if row["detectados"] and row["biometrico"] and row["detectados"] > row["biometrico"] * 1.6:
            findings.append(
                f"{row['hour']}: la camara ve {row['detectados']} personas frente a "
                f"{row['biometrico']:.1f} marcajes (visitas o personal sin marcar).")
        if row["llamadas"] and row["perdidas"] / max(row["llamadas"], 1) > 0.5:
            findings.append(
                f"{row['hour']}: {row['perdidas']:.0f} de {row['llamadas']:.0f} llamadas sin "
                f"atender ({row['perdidas']/max(row['llamadas'],1)*100:.0f}%).")
    peak = max(table, key=lambda r: r["llamadas"]) if table else None
    if peak and peak["llamadas"]:
        findings.insert(0, f"Hora pico de llamadas: {peak['hour']} "
                           f"({peak['llamadas']:.0f} llamadas, {peak['biometrico']:.1f} personas en piso).")
    return {"day": day, "mode": mode, "etiqueta": etiqueta,
            "table": table, "findings": findings[:12],
            "dias_disponibles": common_days(30)}


def business_summary() -> Dict[str, Any]:
    emp = db.query_one("SELECT COUNT(*) n FROM employees") or {"n": 0}
    att = db.query_one(
        """SELECT COUNT(*) n, ROUND(AVG(worked_min)/60.0,2) horas_prom,
                  ROUND(AVG(late_min),1) retraso_prom, MIN(date) desde, MAX(date) hasta
           FROM attendance""") or {}
    calls = db.query_one(
        """SELECT COUNT(*) n, ROUND(AVG(duration_s),1) dur_prom,
                  SUM(CASE WHEN result LIKE '%connected%' THEN 1 ELSE 0 END) atendidas,
                  MIN(date) desde, MAX(date) hasta FROM calls""") or {}
    areas = db.query(
        "SELECT area, COUNT(*) n FROM employees GROUP BY area ORDER BY n DESC")
    top_ext = db.query(
        """SELECT extension, COUNT(*) llamadas,
                  SUM(CASE WHEN result LIKE '%connected%' THEN 1 ELSE 0 END) atendidas
           FROM calls WHERE extension <> '' GROUP BY extension
           ORDER BY llamadas DESC LIMIT 8""")
    puntualidad = db.query(
        """SELECT name, ROUND(AVG(late_min),1) retraso_prom, COUNT(*) dias
           FROM attendance WHERE late_min IS NOT NULL
           GROUP BY emp_id HAVING dias >= 3 ORDER BY retraso_prom DESC LIMIT 8""")
    return {"empleados": emp.get("n", 0), "asistencia": att, "llamadas": calls,
            "areas": areas, "extensiones": top_ext, "puntualidad": puntualidad}
