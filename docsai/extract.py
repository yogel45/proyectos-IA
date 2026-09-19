"""Extraccion de texto de documentos de cualquier origen.

Orden de intentos por formato, siempre degradando sin romper:

  PDF    -> texto nativo con PyMuPDF; si el PDF es un escaneo (apenas hay
            texto) se rasteriza y se pasa por OCR cuando hay Tesseract.
  DOCX   -> parrafos y tablas con python-docx.
  Imagen -> OCR directo (si hay Tesseract).
  EML    -> cabeceras + cuerpo con la libreria estandar.
  TXT/CSV/HTML -> lectura directa con deteccion de codificacion.

Siempre se devuelve la misma estructura, de modo que el resto del flujo
(entidades, clasificacion, nombrado) no sabe de que formato vino el texto.
"""
from __future__ import annotations

import html
import logging
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional

log = logging.getLogger("officevision.docs.extract")

PDF_EXT = {".pdf"}
DOCX_EXT = {".docx", ".dotx"}
IMG_EXT = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}
MAIL_EXT = {".eml"}
TEXT_EXT = {".txt", ".csv", ".md", ".log", ".htm", ".html", ".json", ".xml"}
SOPORTADOS = PDF_EXT | DOCX_EXT | IMG_EXT | MAIL_EXT | TEXT_EXT

# Un PDF con menos de esto por pagina casi seguro es un escaneo
MIN_CHARS_POR_PAGINA = 90


@dataclass
class Extraccion:
    texto: str = ""
    paginas: int = 0
    metodo: str = ""                 # pdf_texto | pdf_ocr | docx | ocr | eml | texto
    idioma_probable: str = ""
    metadatos: Dict[str, Any] = field(default_factory=dict)
    requiere_ocr: bool = False
    ocr_disponible: bool = True
    error: str = ""

    @property
    def caracteres(self) -> int:
        return len(self.texto)

    def resumen(self, limite: int = 400) -> str:
        limpio = re.sub(r"\s+", " ", self.texto).strip()
        return limpio[:limite]


def _tesseract_disponible() -> bool:
    import shutil
    try:
        import pytesseract  # noqa: F401
    except ImportError:
        return False
    return bool(shutil.which("tesseract"))


def _ocr_imagen(ruta_o_bytes) -> str:
    import pytesseract
    from PIL import Image
    import io

    img = (Image.open(io.BytesIO(ruta_o_bytes))
           if isinstance(ruta_o_bytes, (bytes, bytearray))
           else Image.open(str(ruta_o_bytes)))
    return pytesseract.image_to_string(img, lang="spa+eng")


def _detectar_idioma(texto: str) -> str:
    """Heuristica simple por palabras funcionales, suficiente para etiquetar."""
    muestra = texto.lower()[:4000]
    es = sum(muestra.count(p) for p in (" de ", " la ", " que ", " el ", " los ", " para "))
    en = sum(muestra.count(p) for p in (" the ", " of ", " and ", " to ", " for ", " is "))
    if es == en == 0:
        return ""
    return "es" if es >= en else "en"


def extraer(ruta: Path) -> Extraccion:
    ext = ruta.suffix.lower()
    try:
        if ext in PDF_EXT:
            res = _extraer_pdf(ruta)
        elif ext in DOCX_EXT:
            res = _extraer_docx(ruta)
        elif ext in IMG_EXT:
            res = _extraer_imagen(ruta)
        elif ext in MAIL_EXT:
            res = _extraer_eml(ruta)
        elif ext in TEXT_EXT:
            res = _extraer_texto(ruta)
        else:
            return Extraccion(error=f"Formato no soportado: {ext}")
    except Exception as exc:  # pragma: no cover - depende del archivo
        log.exception("Fallo extrayendo %s", ruta.name)
        return Extraccion(error=f"{type(exc).__name__}: {exc}")
    res.idioma_probable = _detectar_idioma(res.texto)
    return res


def _extraer_pdf(ruta: Path) -> Extraccion:
    import pymupdf

    doc = pymupdf.open(str(ruta))
    partes = [pagina.get_text() for pagina in doc]
    texto = "\n".join(partes)
    metadatos = {k: v for k, v in (doc.metadata or {}).items() if v}
    paginas = doc.page_count
    escaneado = paginas > 0 and len(texto.strip()) / max(paginas, 1) < MIN_CHARS_POR_PAGINA

    if escaneado:
        if _tesseract_disponible():
            trozos = []
            for pagina in doc:                      # rasteriza a 200 dpi y aplica OCR
                pix = pagina.get_pixmap(dpi=200)
                trozos.append(_ocr_imagen(pix.tobytes("png")))
            doc.close()
            return Extraccion(texto="\n".join(trozos), paginas=paginas, metodo="pdf_ocr",
                              metadatos=metadatos, requiere_ocr=True)
        doc.close()
        return Extraccion(texto=texto, paginas=paginas, metodo="pdf_texto",
                          metadatos=metadatos, requiere_ocr=True, ocr_disponible=False)
    doc.close()
    return Extraccion(texto=texto, paginas=paginas, metodo="pdf_texto", metadatos=metadatos)


def _extraer_docx(ruta: Path) -> Extraccion:
    import docx

    d = docx.Document(str(ruta))
    partes: List[str] = [p.text for p in d.paragraphs if p.text.strip()]
    for tabla in d.tables:
        for fila in tabla.rows:
            celdas = [c.text.strip() for c in fila.cells if c.text.strip()]
            if celdas:
                partes.append(" | ".join(celdas))
    props = d.core_properties
    metadatos = {k: str(getattr(props, k)) for k in ("author", "title", "subject", "created")
                 if getattr(props, k, None)}
    return Extraccion(texto="\n".join(partes), paginas=max(1, len(partes) // 45),
                      metodo="docx", metadatos=metadatos)


def _extraer_imagen(ruta: Path) -> Extraccion:
    if not _tesseract_disponible():
        return Extraccion(texto="", paginas=1, metodo="ocr", requiere_ocr=True,
                          ocr_disponible=False,
                          error="Imagen sin texto: instala Tesseract para leerla")
    return Extraccion(texto=_ocr_imagen(ruta), paginas=1, metodo="ocr", requiere_ocr=True)


def _extraer_eml(ruta: Path) -> Extraccion:
    import email
    from email import policy

    mensaje = email.message_from_bytes(ruta.read_bytes(), policy=policy.default)
    cabeceras = {k: str(mensaje.get(k, "")) for k in ("From", "To", "Subject", "Date")}
    cuerpo = ""
    if mensaje.is_multipart():
        for parte in mensaje.walk():
            if parte.get_content_type() == "text/plain":
                cuerpo += parte.get_content()
    else:
        cuerpo = mensaje.get_content()
    encabezado = "\n".join(f"{k}: {v}" for k, v in cabeceras.items() if v)
    return Extraccion(texto=f"{encabezado}\n\n{cuerpo}", paginas=1, metodo="eml",
                      metadatos=cabeceras)


def _extraer_texto(ruta: Path) -> Extraccion:
    datos = ruta.read_bytes()
    texto = ""
    for codificacion in ("utf-8", "latin-1", "cp1252"):
        try:
            texto = datos.decode(codificacion)
            break
        except UnicodeDecodeError:
            continue
    if ruta.suffix.lower() in {".htm", ".html"}:
        texto = html.unescape(re.sub(r"<[^>]+>", " ", texto))
    return Extraccion(texto=texto, paginas=max(1, texto.count("\f") + 1), metodo="texto")
